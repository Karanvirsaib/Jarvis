"""Truth-constrained ATS resume tailoring and document export."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
import json
from pathlib import Path
import re
from typing import Any

from core.llm import LLMClient
from core.memory import MemoryManager


STANDARD_SECTIONS = ("SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "PROJECTS", "EDUCATION", "CERTIFICATIONS")


@dataclass
class ResumeDraft:
    name: str
    contact: str
    target_role: str
    company: str = ""
    summary: str = ""
    skills: list[str] = field(default_factory=list)
    experience: list[dict[str, Any]] = field(default_factory=list)
    projects: list[dict[str, Any]] = field(default_factory=list)
    education: list[dict[str, Any]] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    matched_keywords: list[str] = field(default_factory=list)
    gaps: list[str] = field(default_factory=list)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "ResumeDraft":
        def strings(value: Any) -> list[str]:
            return [str(item).strip() for item in value] if isinstance(value, list) else []

        return cls(
            name=str(data.get("name", "")).strip(),
            contact=str(data.get("contact", "")).strip(),
            target_role=str(data.get("target_role", "")).strip(),
            company=str(data.get("company", "")).strip(),
            summary=str(data.get("summary", "")).strip(),
            skills=strings(data.get("skills")),
            experience=[item for item in data.get("experience", []) if isinstance(item, dict)],
            projects=[item for item in data.get("projects", []) if isinstance(item, dict)],
            education=[item for item in data.get("education", []) if isinstance(item, dict)],
            certifications=strings(data.get("certifications")),
            matched_keywords=strings(data.get("matched_keywords")),
            gaps=strings(data.get("gaps")),
        )


class ResumeTailor:
    """Maintains one source resume and one truth-checked tailored draft."""

    def __init__(self, memory: MemoryManager, llm: LLMClient) -> None:
        self.memory = memory
        self.llm = llm
        self.source_path: Path | None = None
        self.source_text = ""
        self.job_description = ""
        self.draft: ResumeDraft | None = None

    def load(self, path: str | Path) -> str:
        candidate = Path(path).expanduser().resolve()
        if not candidate.is_file():
            raise FileNotFoundError(f"Resume file not found: {candidate}")
        if candidate.suffix.casefold() not in {".docx", ".pdf", ".txt"}:
            raise ValueError("Use a .docx, .pdf, or .txt resume file.")
        text = self.extract_text(candidate)
        if len(text.split()) < 40:
            raise ValueError("That file does not contain enough readable resume text.")
        self.source_path = candidate
        self.source_text = text
        self.draft = None
        return f"Resume loaded: {candidate.name} ({len(text.split()):,} words). Now paste the job description using: tailor resume: <job description>"

    @staticmethod
    def extract_text(path: Path) -> str:
        suffix = path.suffix.casefold()
        if suffix == ".txt":
            return path.read_text(encoding="utf-8", errors="replace").strip()
        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("Word resume support requires python-docx.") from exc
            document = Document(path)
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                paragraphs.extend(cell.text for row in table.rows for cell in row.cells if cell.text.strip())
            return "\n".join(paragraphs).strip()
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF resume support requires pypdf.") from exc
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages).strip()

    def tailor(self, job_description: str) -> str:
        if not self.source_text:
            return "Load your current resume first. Use the attachment button or: load resume \"C:\\path\\resume.docx\""
        job_description = job_description.strip()
        if len(job_description.split()) < 25:
            return "Please paste the complete job description so I can identify responsibilities, requirements, and ATS keywords accurately."
        self.job_description = job_description
        prompt = self._prompt(job_description)
        raw = self.llm.ask(
            prompt, self.memory.context(), (), deep_reasoning=True, response_tokens=1800
        )
        data = self._parse_json(raw)
        draft = ResumeDraft.from_dict(data)
        self._validate(draft)
        self.draft = draft
        strengths = ", ".join(draft.matched_keywords[:12]) or "No exact keyword matches were returned"
        gaps = "; ".join(draft.gaps[:6]) or "No material gaps identified"
        return (
            f"ATS DRAFT READY // {draft.target_role or 'Target role'}\n"
            f"Matched evidence-backed keywords: {strengths}.\n"
            f"Honest gaps: {gaps}.\n\n"
            f"{draft.summary}\n\n"
            "Review the draft, then say 'export resume word' or 'export resume pdf'."
        )

    def _prompt(self, job_description: str) -> str:
        return f"""Create a truthful, ATS-friendly resume tailored to the job description.

NON-NEGOTIABLE FACT RULES:
- The SOURCE RESUME is the factual boundary. Never invent employers, titles, dates, education, certifications, tools, metrics, responsibilities, or achievements.
- Saved profile memory may clarify wording but may not supply missing employment dates, employers, education, or achievements.
- Preserve every employer and official job title exactly. Do not turn a Senior Business Analyst role into Senior Data Analyst.
- A personal/practice project must remain under Projects and must never be presented as employer work.
- Do not claim proficiency in a required skill unless the source supports it. Put unsupported requirements in gaps.
- You may reorder, shorten, and rewrite supported material using exact job-description terminology when meaning remains true.
- Use concise accomplishment bullets. Do not create numerical results that are absent from the source.
- Output JSON only, with no markdown or commentary.

JSON SCHEMA:
{{"name":"", "contact":"", "target_role":"", "company":"", "summary":"2-3 lines", "skills":["keyword"], "experience":[{{"title":"", "company":"", "location":"", "dates":"", "bullets":[""]}}], "projects":[{{"name":"", "details":"", "bullets":[""]}}], "education":[{{"credential":"", "institution":"", "dates":""}}], "certifications":[""], "matched_keywords":[""], "gaps":[""]}}

SOURCE RESUME:
---
{self.source_text}
---

JOB DESCRIPTION:
---
{job_description}
---"""

    @staticmethod
    def _parse_json(raw: str) -> dict[str, Any]:
        cleaned = raw.strip()
        if cleaned.startswith("```"):
            cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
        try:
            value = json.loads(cleaned)
        except json.JSONDecodeError as exc:
            raise ValueError("The local model returned an incomplete resume draft. Please retry in Deep mode.") from exc
        if not isinstance(value, dict):
            raise ValueError("The local model returned an invalid resume draft.")
        return value

    def _validate(self, draft: ResumeDraft) -> None:
        if not draft.name or not draft.summary or not draft.experience:
            raise ValueError("The draft is missing essential resume sections; no document was exported.")
        source = self._normalized(self.source_text)
        for role in draft.experience:
            title = str(role.get("title", "")).strip()
            company = str(role.get("company", "")).strip()
            dates = str(role.get("dates", "")).strip()
            for label, value in (("job title", title), ("employer", company), ("employment dates", dates)):
                if value and self._normalized(value) not in source:
                    raise ValueError(f"Truth check blocked an unsupported {label}: {value}. No document was exported.")
        if "senior data analyst" not in source:
            for role in draft.experience:
                if "senior data analyst" in str(role.get("title", "")).casefold():
                    raise ValueError("Truth check blocked Senior Data Analyst as an employment title. Your official title remains unchanged.")

    @staticmethod
    def _normalized(value: str) -> str:
        return " ".join(re.findall(r"[a-z0-9]+", value.casefold()))

    def status(self) -> str:
        source = self.source_path.name if self.source_path else "not loaded"
        draft = "ready" if self.draft else "not created"
        return f"Resume source: {source}. Tailored ATS draft: {draft}."

    def export(self, kind: str, output_root: str | Path = "output") -> Path:
        if not self.draft:
            raise ValueError("Create and review a tailored draft first with: tailor resume: <job description>")
        root = Path(output_root)
        base = self._filename()
        if kind == "word":
            target = root / "resumes" / f"{base}.docx"
            self._write_docx(target)
        elif kind == "pdf":
            target = root / "pdf" / f"{base}.pdf"
            self._write_pdf(target)
        else:
            raise ValueError("Export format must be word or pdf.")
        return target.resolve()

    def _filename(self) -> str:
        assert self.draft is not None
        parts = [self.draft.name, self.draft.company, self.draft.target_role, "Resume"]
        clean = "_".join(filter(None, (re.sub(r"[^A-Za-z0-9]+", "_", item).strip("_") for item in parts)))
        return clean[:120] or f"Tailored_Resume_{datetime.now():%Y%m%d}"

    def _write_docx(self, target: Path) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor

        assert self.draft is not None
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        section = doc.sections[0]
        section.page_height, section.page_width = Inches(11), Inches(8.5)
        section.top_margin = section.bottom_margin = Inches(0.62)
        section.left_margin = section.right_margin = Inches(0.7)
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Arial", Pt(9.5)
        normal.paragraph_format.space_after = Pt(2.5)
        normal.paragraph_format.line_spacing = 1.03

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.draft.name.upper())
        run.bold, run.font.name, run.font.size = True, "Arial", Pt(18)
        contact = doc.add_paragraph(self.draft.contact)
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(7)

        def heading(text: str) -> None:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(text)
            run.bold, run.font.name, run.font.size = True, "Arial", Pt(10.5)
            run.font.color.rgb = RGBColor(31, 58, 78)

        def bullet(text: str) -> None:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.14)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.add_run(text)

        heading("SUMMARY")
        doc.add_paragraph(self.draft.summary)
        heading("SKILLS")
        doc.add_paragraph(" | ".join(self.draft.skills))
        heading("PROFESSIONAL EXPERIENCE")
        for item in self.draft.experience:
            line = " | ".join(filter(None, [str(item.get("title", "")), str(item.get("company", "")), str(item.get("location", ""))]))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.add_run(line).bold = True
            if item.get("dates"):
                p.add_run(f" | {item['dates']}")
            for value in item.get("bullets", []):
                bullet(str(value))
        if self.draft.projects:
            heading("PROJECTS")
            for item in self.draft.projects:
                p = doc.add_paragraph()
                p.add_run(str(item.get("name", ""))).bold = True
                if item.get("details"):
                    p.add_run(f" | {item['details']}")
                for value in item.get("bullets", []):
                    bullet(str(value))
        if self.draft.education:
            heading("EDUCATION")
            for item in self.draft.education:
                doc.add_paragraph(" | ".join(filter(None, [str(item.get("credential", "")), str(item.get("institution", "")), str(item.get("dates", ""))])))
        if self.draft.certifications:
            heading("CERTIFICATIONS")
            doc.add_paragraph(" | ".join(self.draft.certifications))
        doc.save(target)

    def _write_pdf(self, target: Path) -> None:
        try:
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise RuntimeError("PDF resume support requires reportlab.") from exc
        assert self.draft is not None
        target.parent.mkdir(parents=True, exist_ok=True)
        styles = getSampleStyleSheet()
        body = ParagraphStyle("ATSBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=11.2, spaceAfter=3)
        name = ParagraphStyle("ATSName", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=21, alignment=TA_CENTER)
        center = ParagraphStyle("ATSCenter", parent=body, alignment=TA_CENTER, spaceAfter=8)
        section = ParagraphStyle("ATSSection", parent=body, fontName="Helvetica-Bold", fontSize=10.5, textColor="#1F3A4E", spaceBefore=6, spaceAfter=3)
        bold = ParagraphStyle("ATSBold", parent=body, fontName="Helvetica-Bold", spaceBefore=2)
        doc = SimpleDocTemplate(str(target), pagesize=letter, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.62*inch, bottomMargin=.62*inch)
        story = [Paragraph(self._xml(self.draft.name.upper()), name), Paragraph(self._xml(self.draft.contact), center)]
        def add_section(label: str) -> None: story.append(Paragraph(label, section))
        add_section("SUMMARY"); story.append(Paragraph(self._xml(self.draft.summary), body))
        add_section("SKILLS"); story.append(Paragraph(self._xml(" | ".join(self.draft.skills)), body))
        add_section("PROFESSIONAL EXPERIENCE")
        for item in self.draft.experience:
            header = " | ".join(filter(None, [str(item.get("title", "")), str(item.get("company", "")), str(item.get("location", "")), str(item.get("dates", ""))]))
            story.append(Paragraph(self._xml(header), bold))
            for value in item.get("bullets", []): story.append(Paragraph(f"• {self._xml(str(value))}", body))
        if self.draft.projects:
            add_section("PROJECTS")
            for item in self.draft.projects:
                story.append(Paragraph(self._xml(" | ".join(filter(None, [str(item.get("name", "")), str(item.get("details", ""))]))), bold))
                for value in item.get("bullets", []): story.append(Paragraph(f"• {self._xml(str(value))}", body))
        if self.draft.education:
            add_section("EDUCATION")
            for item in self.draft.education: story.append(Paragraph(self._xml(" | ".join(filter(None, [str(item.get("credential", "")), str(item.get("institution", "")), str(item.get("dates", ""))]))), body))
        if self.draft.certifications:
            add_section("CERTIFICATIONS"); story.append(Paragraph(self._xml(" | ".join(self.draft.certifications)), body))
        doc.build(story)

    @staticmethod
    def _xml(value: str) -> str:
        from xml.sax.saxutils import escape
        return escape(value).replace("\n", "<br/>")
