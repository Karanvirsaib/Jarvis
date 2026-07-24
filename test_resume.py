import tempfile
import unittest
from pathlib import Path

from core.memory import MemoryManager
from core.resume import ResumeDraft, ResumeTailor


SOURCE = """Karan Example
karan@example.com | linkedin.com/in/karan
Senior Business Analyst | Acme Financial | Toronto, Canada | Jan 2021 - Present
Translated business requirements into reporting solutions using Excel, SQL, and Power BI.
Built dashboards and documented requirements for operational stakeholders.
Business Analyst | Example Bank | Toronto, Canada | Jun 2018 - Dec 2020
Analyzed customer operations and prepared management reporting in Excel.
Project: Credit Risk Practice Model
Built a personal XGBoost classification project in Python using public sample data.
Education: Bachelor of Commerce | Example University | 2018
Skills: Excel, SQL, Power BI, Tableau, Python, requirements gathering, data analysis
Additional experience includes stakeholder workshops, quality assurance, documentation, dashboard development, and process improvement.
"""


class FakeLLM:
    def __init__(self, output="{}"): self.output = output
    def ask(self, *_args, **_kwargs): return self.output


class ResumeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.memory = MemoryManager(self.root / "memory.db")
        self.tailor = ResumeTailor(self.memory, FakeLLM())
        source = self.root / "resume.txt"
        source.write_text(SOURCE, encoding="utf-8")
        self.tailor.load(source)

    def tearDown(self):
        self.memory.close()
        self.temp.cleanup()

    def test_requires_full_job_description(self):
        self.assertIn("complete job description", self.tailor.tailor("Need analyst with SQL"))

    def test_blocks_invented_employment_title(self):
        draft = ResumeDraft(
            name="Karan Example", contact="karan@example.com", target_role="Senior Data Analyst",
            summary="Evidence-backed candidate.",
            experience=[{"title":"Senior Data Analyst", "company":"Acme Financial", "dates":"Jan 2021 - Present", "bullets":[]}],
        )
        with self.assertRaisesRegex(ValueError, "unsupported job title|Senior Data Analyst"):
            self.tailor._validate(draft)

    def test_exports_single_column_word_and_pdf(self):
        self.tailor.draft = ResumeDraft(
            name="Karan Example", contact="karan@example.com | linkedin.com/in/karan",
            target_role="Data Analyst", company="Example Health", summary="Business analyst with supported reporting and analytics experience.",
            skills=["Excel", "SQL", "Power BI"],
            experience=[{"title":"Senior Business Analyst", "company":"Acme Financial", "location":"Toronto, Canada", "dates":"Jan 2021 - Present", "bullets":["Translated business requirements into reporting solutions using Excel, SQL, and Power BI."]}],
            projects=[{"name":"Credit Risk Practice Model", "details":"Personal project", "bullets":["Built an XGBoost classification project in Python using public sample data."]}],
            education=[{"credential":"Bachelor of Commerce", "institution":"Example University", "dates":"2018"}],
        )
        docx_path = self.tailor.export("word", self.root / "output")
        pdf_path = self.tailor.export("pdf", self.root / "output")
        from docx import Document
        from pypdf import PdfReader
        document = Document(docx_path)
        self.assertEqual(len(document.tables), 0)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("PROFESSIONAL EXPERIENCE", text)
        self.assertIn("Senior Business Analyst", text)
        self.assertIn("SUMMARY", "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages))


if __name__ == "__main__":
    unittest.main()
